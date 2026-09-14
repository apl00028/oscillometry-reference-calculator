from pathlib import Path

from openpyxl import load_workbook

from equations import (
    gochicoa_r5r20,
    gochicoa_x5,
    gochicoa_fres,
    gochicoa_ax,
)

from validation import (
    validate_demographics,
    validate_oscillometry_value,
)


REQUIRED_COLUMNS = [
    "Sexo",
    "Edad",
    "Talla_cm",
    "Peso_kg",
]


OPTIONAL_COLUMNS = [
    "R5_R20",
    "X5",
    "Fres",
    "AX",
]


OUTPUT_COLUMNS = [
    "Registro",
    "IMC",
    "R5_R20_predicho",
    "R5_R20_LLN",
    "R5_R20_ULN",
    "R5_R20_z",
    "R5_R20_estado",
    "X5_predicho",
    "X5_LLN",
    "X5_ULN",
    "X5_z",
    "X5_estado",
    "Fres_predicho",
    "Fres_LLN",
    "Fres_ULN",
    "Fres_z",
    "Fres_estado",
    "AX_predicho",
    "AX_LLN",
    "AX_ULN",
    "AX_z",
    "AX_estado",
    "Advertencias",
    "Error",
]


def status(abnormal):
    return "ALTERADO" if abnormal else "NORMAL"


def empty_parameter_result(prefix):
    return {
        f"{prefix}_predicho": "",
        f"{prefix}_LLN": "",
        f"{prefix}_ULN": "",
        f"{prefix}_z": "",
        f"{prefix}_estado": "NO INTRODUCIDO",
    }


def process_excel(input_path, output_path=None):
    from copy import copy

    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    input_path = Path(input_path)

    if output_path is None:
        output_path = input_path
    else:
        output_path = Path(output_path)

    workbook = load_workbook(input_path)

    # Preferir la hoja creada por nuestra plantilla.
    if "Datos" in workbook.sheetnames:
        source_sheet = workbook["Datos"]
    else:
        source_sheet = workbook.active

    headers = {
        cell.value: cell.column
        for cell in source_sheet[1]
        if cell.value is not None
    }

    missing = [
        column
        for column in REQUIRED_COLUMNS
        if column not in headers
    ]

    if missing:
        raise ValueError(
            "Faltan columnas obligatorias: "
            + ", ".join(missing)
        )

    # Columnas originales, manteniendo su orden.
    source_headers = [
        cell.value
        for cell in source_sheet[1]
        if cell.value is not None
    ]

    # Cada ejecuci?n genera una hoja Resultados limpia.
    results_name = "Resultados"

    if results_name in workbook.sheetnames:
        del workbook[results_name]

    source_position = workbook.index(source_sheet)

    results_sheet = workbook.create_sheet(
        results_name,
        source_position + 1,
    )

    calculated_columns = [
        column
        for column in OUTPUT_COLUMNS
        if column != "Registro"
    ]

    result_headers = (
        ["Registro"]
        + source_headers
        + calculated_columns
    )

    # ---------------------------------------------------------
    # Cabeceras
    # ---------------------------------------------------------

    for column, name in enumerate(
        result_headers,
        start=1,
    ):
        results_sheet.cell(
            row=1,
            column=column,
            value=name,
        )

    # Registro: estilo independiente.
    registro_header = results_sheet["A1"]
    registro_header.fill = PatternFill(
        "solid",
        fgColor="44546A",
    )
    registro_header.font = Font(
        color="FFFFFF",
        bold=True,
    )
    registro_header.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    # Copiar estilo de las cabeceras originales.
    for index, header in enumerate(
        source_headers,
        start=2,
    ):
        source_cell = source_sheet.cell(
            row=1,
            column=headers[header],
        )

        target_cell = results_sheet.cell(
            row=1,
            column=index,
        )

        target_cell._style = copy(
            source_cell._style
        )

    # Columnas transformadas: azul.
    calculated_start = 2 + len(source_headers)

    transformed_header_fill = PatternFill(
        "solid",
        fgColor="2F75B5",
    )

    transformed_fill = PatternFill(
        "solid",
        fgColor="EAF3F8",
    )

    for column in range(
        calculated_start,
        len(result_headers) + 1,
    ):
        cell = results_sheet.cell(
            row=1,
            column=column,
        )

        cell.fill = transformed_header_fill
        cell.font = Font(
            color="FFFFFF",
            bold=True,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # ---------------------------------------------------------
    # Procesamiento
    # ---------------------------------------------------------

    processed = 0
    errors = 0
    record_number = 0

    def get_value(row, column_name):
        column = headers.get(column_name)

        if column is None:
            return None

        return source_sheet.cell(
            row=row,
            column=column,
        ).value

    source_columns = [
        column
        for column in REQUIRED_COLUMNS + OPTIONAL_COLUMNS
        if column in headers
    ]

    for source_row in range(
        2,
        source_sheet.max_row + 1,
    ):
        row_values = [
            get_value(source_row, column)
            for column in source_columns
        ]

        if all(
            value is None or str(value).strip() == ""
            for value in row_values
        ):
            continue

        record_number += 1
        registro = f"P{record_number:03d}"

        target_row = results_sheet.max_row + 1

        # ID siempre en primera columna.
        results_sheet.cell(
            row=target_row,
            column=1,
            value=registro,
        )

        # Copiar los datos originales a continuaci?n.
        for offset, header in enumerate(
            source_headers,
            start=2,
        ):
            value = get_value(
                source_row,
                header,
            )

            results_sheet.cell(
                row=target_row,
                column=offset,
                value=value,
            )

        values = {
            name: ""
            for name in calculated_columns
        }

        try:
            sex = get_value(
                source_row,
                "Sexo",
            )

            demographics = validate_demographics(
                sex=sex,
                age=get_value(source_row, "Edad"),
                height_cm=get_value(source_row, "Talla_cm"),
                weight_kg=get_value(source_row, "Peso_kg"),
            )

            values["Advertencias"] = " | ".join(
                demographics["warnings"]
            )

            if demographics["errors"]:
                raise ValueError(
                    " | ".join(
                        demographics["errors"]
                    )
                )

            age = demographics["age"]
            height = demographics["height_cm"]
            weight = demographics["weight_kg"]

            values["IMC"] = demographics["bmi"]

            r5r20_value = validate_oscillometry_value(
                "R5_R20",
                get_value(source_row, "R5_R20"),
            )

            x5_value = validate_oscillometry_value(
                "X5",
                get_value(source_row, "X5"),
            )

            fres_value = validate_oscillometry_value(
                "Fres",
                get_value(source_row, "Fres"),
            )

            ax_value = validate_oscillometry_value(
                "AX",
                get_value(source_row, "AX"),
            )

            if r5r20_value is None:
                values.update(
                    empty_parameter_result(
                        "R5_R20"
                    )
                )
            else:
                result = gochicoa_r5r20(
                    age,
                    sex,
                    height,
                    weight,
                    r5r20_value,
                )

                values.update({
                    "R5_R20_predicho": result["predicted"],
                    "R5_R20_LLN": result["lln"],
                    "R5_R20_ULN": result["uln"],
                    "R5_R20_z": result["z_score"],
                    "R5_R20_estado": status(
                        result["abnormal"]
                    ),
                })

            if x5_value is None:
                values.update(
                    empty_parameter_result("X5")
                )
            else:
                result = gochicoa_x5(
                    age,
                    sex,
                    height,
                    weight,
                    x5_value,
                )

                values.update({
                    "X5_predicho": result["predicted"],
                    "X5_LLN": result["lln"],
                    "X5_ULN": result["uln"],
                    "X5_z": result["z_score"],
                    "X5_estado": status(
                        result["abnormal"]
                    ),
                })

            if fres_value is None:
                values.update(
                    empty_parameter_result("Fres")
                )
            else:
                result = gochicoa_fres(
                    age,
                    sex,
                    height,
                    weight,
                    fres_value,
                )

                values.update({
                    "Fres_predicho": result["predicted"],
                    "Fres_LLN": result["lln"],
                    "Fres_ULN": result["uln"],
                    "Fres_z": result["z_score"],
                    "Fres_estado": status(
                        result["abnormal"]
                    ),
                })

            if ax_value is None:
                values.update(
                    empty_parameter_result("AX")
                )
            else:
                result = gochicoa_ax(
                    age,
                    sex,
                    height,
                    weight,
                    ax_value,
                )

                values.update({
                    "AX_predicho": result["predicted"],
                    "AX_LLN": result["lln"],
                    "AX_ULN": result["uln"],
                    "AX_z": result["z_score"],
                    "AX_estado": status(
                        result["abnormal"]
                    ),
                })

            processed += 1

        except (ValueError, TypeError) as exc:
            values["Error"] = str(exc)
            errors += 1

        # Escribir columnas calculadas.
        for offset, name in enumerate(
            calculated_columns,
            start=calculated_start,
        ):
            cell = results_sheet.cell(
                row=target_row,
                column=offset,
                value=values[name],
            )

            cell.fill = transformed_fill

        # Registro en gris suave.
        results_sheet.cell(
            row=target_row,
            column=1,
        ).fill = PatternFill(
            "solid",
            fgColor="E7E6E6",
        )

    # ---------------------------------------------------------
    # Formato visual
    # ---------------------------------------------------------

    results_sheet.freeze_panes = "A2"
    results_sheet.row_dimensions[1].height = 30

    results_sheet.column_dimensions["A"].width = 12

    # Copiar anchos aproximados de las columnas originales.
    for offset, header in enumerate(
        source_headers,
        start=2,
    ):
        source_column = get_column_letter(
            headers[header]
        )
        target_column = get_column_letter(
            offset
        )

        width = source_sheet.column_dimensions[
            source_column
        ].width

        results_sheet.column_dimensions[
            target_column
        ].width = width or 15

    # Anchos para resultados.
    for column in range(
        calculated_start,
        len(result_headers) + 1,
    ):
        letter = get_column_letter(column)
        results_sheet.column_dimensions[
            letter
        ].width = 18

    # Estados con color propio.
    normal_fill = PatternFill(
        "solid",
        fgColor="D1FADF",
    )
    altered_fill = PatternFill(
        "solid",
        fgColor="FEE4E2",
    )
    missing_fill = PatternFill(
        "solid",
        fgColor="F2F4F7",
    )
    warning_fill = PatternFill(
        "solid",
        fgColor="FFF4CC",
    )
    error_fill = PatternFill(
        "solid",
        fgColor="FDE9E7",
    )

    header_positions = {
        cell.value: cell.column
        for cell in results_sheet[1]
    }

    status_columns = [
        "R5_R20_estado",
        "X5_estado",
        "Fres_estado",
        "AX_estado",
    ]

    for row in range(
        2,
        results_sheet.max_row + 1,
    ):
        for name in status_columns:
            column = header_positions[name]
            cell = results_sheet.cell(
                row=row,
                column=column,
            )

            if cell.value == "NORMAL":
                cell.fill = normal_fill
            elif cell.value == "ALTERADO":
                cell.fill = altered_fill
            elif cell.value == "NO INTRODUCIDO":
                cell.fill = missing_fill

        warning_cell = results_sheet.cell(
            row=row,
            column=header_positions["Advertencias"],
        )

        if warning_cell.value:
            warning_cell.fill = warning_fill

        error_cell = results_sheet.cell(
            row=row,
            column=header_positions["Error"],
        )

        if error_cell.value:
            error_cell.fill = error_fill

    if results_sheet.max_row >= 1:
        last_column = get_column_letter(
            len(result_headers)
        )

        results_sheet.auto_filter.ref = (
            f"A1:{last_column}"
            f"{results_sheet.max_row}"
        )

    # Abrir el libro mostrando directamente Resultados.
    workbook.active = results_sheet

    workbook.save(output_path)

    return {
        "processed": processed,
        "errors": errors,
        "output_path": str(output_path),
        "results_sheet": results_name,
    }


def create_template(output_path):
    from openpyxl import Workbook
    from openpyxl.styles import (
        Alignment,
        Border,
        Font,
        PatternFill,
        Side,
    )
    from openpyxl.worksheet.datavalidation import DataValidation

    output_path = Path(output_path)

    workbook = Workbook()

    data_sheet = workbook.active
    data_sheet.title = "Datos"

    input_columns = [
        "Sexo",
        "Edad",
        "Talla_cm",
        "Peso_kg",
        "R5_R20",
        "X5",
        "Fres",
        "AX",
    ]

    # ---------------------------------------------------------
    # Estilo general
    # ---------------------------------------------------------

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    required_fill = PatternFill(
        "solid",
        fgColor="EAF2F8",
    )

    optional_fill = PatternFill(
        "solid",
        fgColor="FFFBEA",
    )

    section_fill = PatternFill(
        "solid",
        fgColor="D9EAF7",
    )

    thin_gray = Side(
        style="thin",
        color="D0D5DD",
    )

    cell_border = Border(
        bottom=thin_gray,
    )

    # ---------------------------------------------------------
    # Hoja Datos
    # ---------------------------------------------------------

    for column, name in enumerate(
        input_columns,
        start=1,
    ):
        cell = data_sheet.cell(
            row=1,
            column=column,
            value=name,
        )

        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    data_sheet.row_dimensions[1].height = 30
    data_sheet.freeze_panes = "A2"
    data_sheet.auto_filter.ref = "A1:H501"

    widths = {
        "A": 16,
        "B": 13,
        "C": 15,
        "D": 15,
        "E": 17,
        "F": 15,
        "G": 15,
        "H": 15,
    }

    for column, width in widths.items():
        data_sheet.column_dimensions[column].width = width

    # Diferenciar campos obligatorios y opcionales en la cabecera
    for column in range(1, 9):
        cell = data_sheet.cell(
            row=1,
            column=column,
        )

        if column <= 4:
            cell.fill = PatternFill(
                "solid",
                fgColor="1F4E78",
            )
        else:
            cell.fill = PatternFill(
                "solid",
                fgColor="B7791F",
            )

        cell.font = header_font

    # ---------------------------------------------------------
    # Validaciones de entrada
    # ---------------------------------------------------------

    sex_validation = DataValidation(
        type="list",
        formula1='"Mujer,Hombre"',
        allow_blank=False,
    )

    sex_validation.error = (
        "Seleccione Mujer u Hombre."
    )
    sex_validation.errorTitle = "Sexo no v?lido"
    sex_validation.prompt = (
        "Seleccione el sexo en el desplegable."
    )
    sex_validation.promptTitle = "Sexo"

    data_sheet.add_data_validation(
        sex_validation
    )
    sex_validation.add("A2:A501")

    age_validation = DataValidation(
        type="decimal",
        operator="between",
        formula1="2.7",
        formula2="90",
        allow_blank=False,
    )

    age_validation.error = (
        "La edad debe estar entre 2,7 y 90 a?os."
    )
    age_validation.errorTitle = "Edad fuera de rango"
    age_validation.prompt = (
        "Rango publicado de Gochicoa-Rangel: 2,7-90 a?os."
    )
    age_validation.promptTitle = "Edad"

    data_sheet.add_data_validation(
        age_validation
    )
    age_validation.add("B2:B501")

    height_validation = DataValidation(
        type="decimal",
        operator="greaterThan",
        formula1="0",
        allow_blank=False,
    )

    height_validation.error = (
        "Introduzca la talla en cent?metros y con un valor mayor que 0."
    )
    height_validation.errorTitle = "Talla no v?lida"
    height_validation.prompt = (
        "Introduzca cent?metros, por ejemplo 165 y no 1,65."
    )
    height_validation.promptTitle = "Talla en cm"

    data_sheet.add_data_validation(
        height_validation
    )
    height_validation.add("C2:C501")

    weight_validation = DataValidation(
        type="decimal",
        operator="greaterThan",
        formula1="0",
        allow_blank=False,
    )

    weight_validation.error = (
        "El peso debe ser mayor que 0 kg."
    )
    weight_validation.errorTitle = "Peso no v?lido"
    weight_validation.prompt = (
        "Introduzca el peso en kilogramos."
    )
    weight_validation.promptTitle = "Peso en kg"

    data_sheet.add_data_validation(
        weight_validation
    )
    weight_validation.add("D2:D501")

    fres_validation = DataValidation(
        type="decimal",
        operator="greaterThan",
        formula1="0",
        allow_blank=True,
    )

    fres_validation.error = (
        "Fres debe ser mayor que 0."
    )
    fres_validation.errorTitle = "Fres no v?lida"

    data_sheet.add_data_validation(
        fres_validation
    )
    fres_validation.add("G2:G501")

    ax_validation = DataValidation(
        type="decimal",
        operator="greaterThan",
        formula1="0",
        allow_blank=True,
    )

    ax_validation.error = (
        "AX debe ser mayor que 0."
    )
    ax_validation.errorTitle = "AX no v?lida"

    data_sheet.add_data_validation(
        ax_validation
    )
    ax_validation.add("H2:H501")

    # ---------------------------------------------------------
    # Formatos num?ricos
    # ---------------------------------------------------------


    # ---------------------------------------------------------
    # Hoja Instrucciones
    # ---------------------------------------------------------

    instructions = workbook.create_sheet(
        "Instrucciones"
    )

    instructions.merge_cells("A1:B1")

    title = instructions["A1"]
    title.value = (
        "Calculadora de oscilometr?a ? plantilla de entrada"
    )
    title.font = Font(
        size=16,
        bold=True,
        color="1F4E78",
    )

    instructions.merge_cells("A2:B2")
    instructions["A2"] = (
        "Complete una fila por paciente. "
        "No cambie los nombres de las columnas de la hoja Datos."
    )
    instructions["A2"].font = Font(
        italic=True,
        color="667085",
    )

    instructions.merge_cells("A3:B3")
    instructions["A3"] = (
        "Azul = campos obligatorios. "
        "Amarillo = par?metros opcionales; puede dejar los que no tenga en blanco."
    )
    instructions["A3"].font = Font(
        italic=True,
        color="667085",
    )

    instructions["A5"] = "Campo"
    instructions["B5"] = "Descripci?n"

    for cell in instructions[5]:
        cell.fill = header_fill
        cell.font = header_font

    descriptions = [
        (
            "Sexo",
            "Obligatorio. Seleccione Mujer u Hombre.",
        ),
        (
            "Edad",
            "Obligatorio. Edad en a?os. "
            "Rango de referencia publicado: 2,7-90 a?os.",
        ),
        (
            "Talla_cm",
            "Obligatorio. Talla en cent?metros. "
            "Ejemplo: 165, no 1,65.",
        ),
        (
            "Peso_kg",
            "Obligatorio. Peso en kilogramos.",
        ),
        (
            "R5_R20",
            "Opcional. R5-R20 observado antes de broncodilatador.",
        ),
        (
            "X5",
            "Opcional. X5 observado antes de broncodilatador.",
        ),
        (
            "Fres",
            "Opcional. Frecuencia de resonancia observada "
            "antes de broncodilatador.",
        ),
        (
            "AX",
            "Opcional. ?rea de reactancia observada "
            "antes de broncodilatador. Debe ser mayor que 0.",
        ),
    ]

    for row, (field, description) in enumerate(
        descriptions,
        start=6,
    ):
        field_cell = instructions.cell(
            row=row,
            column=1,
            value=field,
        )

        description_cell = instructions.cell(
            row=row,
            column=2,
            value=description,
        )

        field_cell.font = Font(
            bold=True
        )
        field_cell.fill = section_fill
        field_cell.alignment = Alignment(
            vertical="top"
        )

        description_cell.alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )

    note_row = 16

    instructions.merge_cells(
        start_row=note_row,
        start_column=1,
        end_row=note_row,
        end_column=2,
    )

    note = instructions.cell(
        row=note_row,
        column=1,
    )

    note.value = (
        "La herramienta generar? autom?ticamente un identificador "
        "P001, P002, P003... en el archivo de resultados."
    )

    note.font = Font(
        size=10,
        italic=True,
        color="667085",
    )

    instructions.merge_cells(
        start_row=note_row + 2,
        start_column=1,
        end_row=note_row + 2,
        end_column=2,
    )

    reference = instructions.cell(
        row=note_row + 2,
        column=1,
    )

    reference.value = (
        "Referencia: Gochicoa-Rangel et al. "
        "ERJ Open Research 2023;9:00503-2023."
    )

    reference.font = Font(
        size=10,
        italic=True,
        color="667085",
    )

    instructions.merge_cells(
        start_row=note_row + 4,
        start_column=1,
        end_row=note_row + 4,
        end_column=2,
    )

    disclaimer = instructions.cell(
        row=note_row + 4,
        column=1,
    )

    disclaimer.value = (
        "Herramienta destinada al c?lculo de valores de referencia "
        "y apoyo a investigaci?n. Los resultados deben interpretarse "
        "en su contexto cl?nico."
    )

    disclaimer.font = Font(
        size=10,
        color="667085",
    )

    disclaimer.alignment = Alignment(
        wrap_text=True
    )

    instructions.column_dimensions["A"].width = 20
    instructions.column_dimensions["B"].width = 82

    instructions.freeze_panes = "A6"

    workbook.save(output_path)

    return str(output_path)
