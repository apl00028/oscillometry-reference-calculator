import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

tool_dir = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(tool_dir))

from excel_processor import (
    OUTPUT_COLUMNS,
    create_template,
    process_excel,
)


class TestExcelProcessor(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def create_input(self, rows):
        path = self.folder / "cohorte.xlsx"
        create_template(path)

        wb = load_workbook(path)
        ws = wb["Datos"]

        for row in rows:
            ws.append(row)

        wb.save(path)
        return path

    def test_template_has_no_phantom_rows(self):
        path = self.folder / "template.xlsx"
        create_template(path)

        wb = load_workbook(path)
        ws = wb["Datos"]

        self.assertEqual(ws.max_row, 1)

    def test_processes_same_file_and_creates_results_sheet(self):
        path = self.create_input([
            [
                "Mujer", 60, 165, 65,
                0.08, -0.15, 20, 1.0,
            ],
        ])

        result = process_excel(path)

        self.assertEqual(
            Path(result["output_path"]),
            path,
        )

        wb = load_workbook(path)

        self.assertIn("Datos", wb.sheetnames)
        self.assertIn("Resultados", wb.sheetnames)

        # La hoja original permanece intacta.
        self.assertEqual(
            wb["Datos"]["A1"].value,
            "Sexo",
        )
        self.assertNotEqual(
            wb["Datos"]["A1"].value,
            "Registro",
        )

    def test_registro_is_first_column(self):
        path = self.create_input([
            [
                "Mujer", 60, 165, 65,
                0.08, -0.15, 20, 1.0,
            ],
        ])

        process_excel(path)

        wb = load_workbook(path, data_only=True)
        ws = wb["Resultados"]

        self.assertEqual(
            ws["A1"].value,
            "Registro",
        )

        self.assertEqual(
            ws["A2"].value,
            "P001",
        )

        self.assertEqual(
            ws["B1"].value,
            "Sexo",
        )

    def test_complete_patient_values(self):
        path = self.create_input([
            [
                "Mujer", 60, 165, 65,
                0.08, -0.15, 20, 1.0,
            ],
        ])

        result = process_excel(path)

        self.assertEqual(result["processed"], 1)
        self.assertEqual(result["errors"], 0)

        wb = load_workbook(path, data_only=True)
        ws = wb["Resultados"]

        headers = {
            cell.value: cell.column
            for cell in ws[1]
        }

        self.assertAlmostEqual(
            ws.cell(
                2,
                headers["R5_R20_ULN"],
            ).value,
            0.11475622675291378,
        )

        self.assertAlmostEqual(
            ws.cell(
                2,
                headers["X5_LLN"],
            ).value,
            -0.17139451396503497,
        )

        self.assertAlmostEqual(
            ws.cell(
                2,
                headers["AX_ULN"],
            ).value,
            1.3684543809200072,
        )

    def test_transformed_columns_have_colour(self):
        path = self.create_input([
            [
                "Mujer", 60, 165, 65,
                0.08, -0.15, 20, 1.0,
            ],
        ])

        process_excel(path)

        wb = load_workbook(path)
        ws = wb["Resultados"]

        headers = {
            cell.value: cell.column
            for cell in ws[1]
        }

        imc_column = headers["IMC"]

        header_colour = (
            ws.cell(
                1,
                imc_column,
            ).fill.fgColor.rgb
            or ""
        )

        value_colour = (
            ws.cell(
                2,
                imc_column,
            ).fill.fgColor.rgb
            or ""
        )

        self.assertTrue(
            header_colour.endswith("2F75B5")
        )

        self.assertTrue(
            value_colour.endswith("EAF3F8")
        )

    def test_missing_oscillometry_is_allowed(self):
        path = self.create_input([
            [
                "Hombre", 55, 178, 80,
                0.09, None, None, None,
            ],
        ])

        process_excel(path)

        wb = load_workbook(path, data_only=True)
        ws = wb["Resultados"]

        headers = {
            cell.value: cell.column
            for cell in ws[1]
        }

        self.assertEqual(
            ws.cell(
                2,
                headers["AX_estado"],
            ).value,
            "NO INTRODUCIDO",
        )

    def test_height_in_metres_generates_error(self):
        path = self.create_input([
            [
                "Mujer", 60, 1.65, 65,
                0.08, -0.15, 20, 1.0,
            ],
        ])

        result = process_excel(path)

        self.assertEqual(result["processed"], 0)
        self.assertEqual(result["errors"], 1)

        wb = load_workbook(path, data_only=True)
        ws = wb["Resultados"]

        headers = {
            cell.value: cell.column
            for cell in ws[1]
        }

        error = ws.cell(
            2,
            headers["Error"],
        ).value

        self.assertIn("metros", error)

    def test_second_run_replaces_results_sheet(self):
        path = self.create_input([
            [
                "Mujer", 60, 165, 65,
                0.08, -0.15, 20, 1.0,
            ],
        ])

        process_excel(path)
        process_excel(path)

        wb = load_workbook(path)

        self.assertEqual(
            wb.sheetnames.count("Resultados"),
            1,
        )


if __name__ == "__main__":
    unittest.main()
